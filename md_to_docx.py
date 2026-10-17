import docx
from docx.shared import Pt, RGBColor
import re

def convert_md_to_docx(md_filepath, docx_filepath):
    doc = docx.Document()
    
    with open(md_filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            p = doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            p = doc.add_heading(line[4:], level=3)
        elif line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            _add_formatted_text(p, line[2:])
        elif re.match(r'^\d+\.\s', line):
            # Numbered list
            p = doc.add_paragraph(style='List Number')
            text_content = re.sub(r'^\d+\.\s', '', line)
            _add_formatted_text(p, text_content)
        else:
            p = doc.add_paragraph()
            _add_formatted_text(p, line)
            
    doc.save(docx_filepath)

def _add_formatted_text(paragraph, text):
    # Simple bold parser: **text**
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            # Simple italic parser: *text* (if not inside bold)
            subparts = re.split(r'(\*[^\*]+\*)', part)
            for subpart in subparts:
                if subpart.startswith('*') and subpart.endswith('*'):
                    run = paragraph.add_run(subpart[1:-1])
                    run.italic = True
                else:
                    paragraph.add_run(subpart)

if __name__ == '__main__':
    convert_md_to_docx(r'c:\Users\HP\OneDrive\Desktop\mini-qunatum\docs\COMPREHENSIVE_REVIEW_GUIDE.md', r'c:\Users\HP\OneDrive\Desktop\mini-qunatum\docs\Project_Review_Master_Guide.docx')
    print("Done generating DOCX")
